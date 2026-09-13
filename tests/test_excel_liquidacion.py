"""Tests de extractores/excel_liquidacion.py contra .xlsx sintéticos.

Los archivos reales (CAQUETA 01.09.26.xlsx, etc.) NO se copian al repo -- es
público -- así que acá se generan con openpyxl en tmp_path, replicando la
estructura verificada contra esos 3 archivos: arranque en columna B, título
combinado por encima de los encabezados, fórmulas '=precio*cantidad' para el
importe de línea y '=SUM(...)' para el total (openpyxl no cachea el
resultado de una fórmula que él mismo escribe -- por eso casi todos los
casos de acá ejercitan justo la vía de "fórmula sin valor cacheado", que es
la normal para estos tests y también el caso real más común).
"""
from __future__ import annotations

import pathlib
import sys
from datetime import date

import openpyxl
import pytest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

from extractores import excel_liquidacion

_ENCABEZADOS = ["CANTIDAD", "U. DE MEDIDA", "PRODUCTO", "P. UNITARIO", "IMPORTE "]


def _crear_xlsx(
    tmp_path: pathlib.Path,
    nombre_archivo: str,
    filas: list[tuple[float | None, str | None, str, float | None]],
    *,
    titulo: str | None = "VERDURA CAQUETA 05.09.26",
    con_encabezados: bool = True,
    con_total: bool = True,
    fila_titulo: int = 2,
    fila_encabezado: int = 3,
    col_inicio: int = 2,  # columna B, como los archivos reales
) -> pathlib.Path:
    """Arma un .xlsx con la misma forma que los archivos reales: título
    combinado, encabezados encontrados por texto (no por posición fija) y
    fórmulas de Excel para importe de línea y total."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja1"

    if titulo is not None:
        ws.cell(row=fila_titulo, column=col_inicio, value=titulo)
        ws.merge_cells(
            start_row=fila_titulo, start_column=col_inicio, end_row=fila_titulo, end_column=col_inicio + 4
        )

    if con_encabezados:
        for i, texto in enumerate(_ENCABEZADOS):
            ws.cell(row=fila_encabezado, column=col_inicio + i, value=texto)

    col_cantidad, col_unidad, col_producto, col_precio, col_importe = (col_inicio + i for i in range(5))

    fila = fila_encabezado + 1
    primera_fila_datos = fila
    for cantidad, unidad, producto, precio in filas:
        if cantidad is not None:
            ws.cell(row=fila, column=col_cantidad, value=cantidad)
        if unidad is not None:
            ws.cell(row=fila, column=col_unidad, value=unidad)
        ws.cell(row=fila, column=col_producto, value=producto)
        if precio is not None:
            ws.cell(row=fila, column=col_precio, value=precio)
        celda_precio = ws.cell(row=fila, column=col_precio).coordinate
        celda_cantidad = ws.cell(row=fila, column=col_cantidad).coordinate
        ws.cell(row=fila, column=col_importe, value=f"={celda_precio}*{celda_cantidad}")
        fila += 1
    ultima_fila_datos = fila - 1

    if con_total:
        ws.cell(row=fila, column=col_inicio, value="TOTAL")
        ws.merge_cells(start_row=fila, start_column=col_inicio, end_row=fila, end_column=col_inicio + 3)
        celda_inicio = ws.cell(row=primera_fila_datos, column=col_importe).coordinate
        celda_fin = ws.cell(row=ultima_fila_datos, column=col_importe).coordinate
        ws.cell(row=fila, column=col_importe, value=f"=SUM({celda_inicio}:{celda_fin})")

    ruta = tmp_path / nombre_archivo
    wb.save(ruta)
    return ruta


# -----------------------------------------------------------------------------
# Caso feliz
# -----------------------------------------------------------------------------
def test_caso_feliz(tmp_path):
    filas = [
        (10, "KG", "LIMON", 1.5),
        (2, "KG", "AJO ", 3.25),
        (5, "PQTE", "CEBOLLA", 2.0),
        (3, "UNIDAD", "HUEVOS", 1.0),
    ]
    ruta = _crear_xlsx(tmp_path, "CAQUETA 05.09.26.xlsx", filas, titulo="VERDURA CAQUETA 05.09.26")

    comp = excel_liquidacion.extraer(ruta)

    assert comp.origen == "excel"
    assert comp.confianza == 1.0
    assert comp.tipo_documento == "liquidacion"
    assert comp.moneda == "PEN"
    assert comp.condicion == "contado"
    assert comp.proveedor_ruc is None
    assert comp.serie_numero is None
    assert comp.cliente_ruc is None
    assert comp.fecha_emision == "2026-09-05"
    assert comp.proveedor_razon_social == "VERDURA CAQUETA"
    assert len(comp.items) == 4
    assert [item.unidad for item in comp.items] == ["kg", "kg", "paq", "unid"]
    assert [item.orden for item in comp.items] == [1, 2, 3, 4]
    total_esperado = round(10 * 1.5 + 2 * 3.25 + 5 * 2.0 + 3 * 1.0, 2)
    assert comp.total == total_esperado
    assert comp.validar(hoy=date(2026, 9, 13)) == []


# -----------------------------------------------------------------------------
# Fórmulas sin valor cacheado (el caso normal con openpyxl) -> se calcula
# -----------------------------------------------------------------------------
def test_formulas_sin_cache_calcula_importes_y_total(tmp_path):
    filas = [(4, "KG", "TOMATE", 2.5), (1, "KG", "CULANTRO", 4.0)]
    ruta = _crear_xlsx(tmp_path, "CAQUETA 09.09.26.xlsx", filas, titulo="VERDURA CAQUETA 09.09.26")

    comp = excel_liquidacion.extraer(ruta)

    assert [item.total_linea for item in comp.items] == [10.0, 4.0]
    assert comp.total == 14.0
    assert not any("no coincide" in a for a in comp.advertencias)


# -----------------------------------------------------------------------------
# Título con año distinto al del nombre del archivo -> gana el nombre,
# se advierte, y NO se lanza excepción (caso real: CAQUETA 01.09.26.xlsx con
# título "VERDURA CAQUETA 01.09.23").
# -----------------------------------------------------------------------------
def test_anio_distinto_entre_nombre_y_titulo_usa_nombre_y_advierte(tmp_path):
    filas = [(1, "KG", "ZANAHORIA", 2.0)]
    ruta = _crear_xlsx(tmp_path, "CAQUETA 01.09.26.xlsx", filas, titulo="VERDURA CAQUETA 01.09.23")

    comp = excel_liquidacion.extraer(ruta)

    assert comp.fecha_emision == "2026-09-01"
    assert comp.proveedor_razon_social == "VERDURA CAQUETA"
    assert any("difiere" in a for a in comp.advertencias)


# -----------------------------------------------------------------------------
# Unidades
# -----------------------------------------------------------------------------
def test_unidad_pque_normaliza_a_paq(tmp_path):
    filas = [(3, "PQUE", "CHOCLO", 1.0)]
    ruta = _crear_xlsx(tmp_path, "CAQUETA 05.09.26.xlsx", filas)

    comp = excel_liquidacion.extraer(ruta)

    assert comp.items[0].unidad == "paq"


def test_unidad_desconocida_da_none_y_advertencia(tmp_path):
    filas = [(3, "SACO", "PAPA", 1.0)]
    ruta = _crear_xlsx(tmp_path, "CAQUETA 05.09.26.xlsx", filas)

    comp = excel_liquidacion.extraer(ruta)

    assert comp.items[0].unidad is None
    assert any("SACO" in a for a in comp.advertencias)


# -----------------------------------------------------------------------------
# Encabezado desplazado (columna A, fila 5 en vez de B, fila 3)
# -----------------------------------------------------------------------------
def test_encabezado_desplazado_se_lee_igual(tmp_path):
    filas = [(2, "KG", "PEPINO", 3.0)]
    ruta = _crear_xlsx(
        tmp_path,
        "CAQUETA 05.09.26.xlsx",
        filas,
        titulo="VERDURA CAQUETA 05.09.26",
        fila_titulo=2,
        fila_encabezado=5,
        col_inicio=1,
    )

    comp = excel_liquidacion.extraer(ruta)

    assert len(comp.items) == 1
    assert comp.items[0].descripcion == "PEPINO"
    assert comp.items[0].total_linea == 6.0


# -----------------------------------------------------------------------------
# Sin encabezados reconocibles -> ValueError
# -----------------------------------------------------------------------------
def test_sin_encabezados_lanza_valueerror(tmp_path):
    ruta = _crear_xlsx(tmp_path, "CAQUETA 05.09.26.xlsx", [(1, "KG", "APIO", 1.0)], con_encabezados=False)

    with pytest.raises(ValueError, match="CANTIDAD/PRODUCTO/IMPORTE"):
        excel_liquidacion.extraer(ruta)


# -----------------------------------------------------------------------------
# Sin fecha ni en el nombre ni en el título -> ValueError
# -----------------------------------------------------------------------------
def test_sin_fecha_en_nombre_ni_titulo_lanza_valueerror(tmp_path):
    ruta = _crear_xlsx(
        tmp_path, "LIQUIDACION CAQUETA.xlsx", [(1, "KG", "RABANITO", 1.0)], titulo="VERDURA CAQUETA"
    )

    with pytest.raises(ValueError):
        excel_liquidacion.extraer(ruta)


# -----------------------------------------------------------------------------
# Sin fila TOTAL -> suma de importes + advertencia
# -----------------------------------------------------------------------------
def test_sin_fila_total_usa_suma_y_advierte(tmp_path):
    filas = [(2, "KG", "BETARRAGA", 2.0), (1, "KG", "NABO", 1.5)]
    ruta = _crear_xlsx(tmp_path, "CAQUETA 05.09.26.xlsx", filas, con_total=False)

    comp = excel_liquidacion.extraer(ruta)

    assert comp.total == round(2 * 2.0 + 1 * 1.5, 2)
    assert any("no se encontró la fila TOTAL" in a for a in comp.advertencias)


# -----------------------------------------------------------------------------
# Producto repetido (LIMON dos veces con precio distinto, caso real) -> 2 ítems
# -----------------------------------------------------------------------------
def test_producto_repetido_da_dos_items(tmp_path):
    filas = [(5, "KG", "LIMON", 1.0), (3, "KG", "LIMON", 1.5)]
    ruta = _crear_xlsx(tmp_path, "CAQUETA 05.09.26.xlsx", filas)

    comp = excel_liquidacion.extraer(ruta)

    assert len(comp.items) == 2
    assert comp.items[0].descripcion == comp.items[1].descripcion == "LIMON"
    assert comp.items[0].total_linea == 5.0
    assert comp.items[1].total_linea == 4.5


# -----------------------------------------------------------------------------
# Fila con producto pero sin cantidad o sin precio -> se agrega igual, con
# advertencia
# -----------------------------------------------------------------------------
def test_fila_sin_cantidad_o_precio_se_agrega_con_advertencia(tmp_path):
    filas = [(None, "KG", "YUCA", 2.0), (4, "KG", "CAMOTE", None)]
    ruta = _crear_xlsx(tmp_path, "CAQUETA 05.09.26.xlsx", filas)

    comp = excel_liquidacion.extraer(ruta)

    assert len(comp.items) == 2
    assert comp.items[0].cantidad is None
    assert comp.items[1].precio_unitario is None
    mensajes = " ".join(comp.advertencias)
    assert "YUCA" in mensajes and "CAMOTE" in mensajes
